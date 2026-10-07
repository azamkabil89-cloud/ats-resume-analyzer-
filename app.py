import json
import os
import re
from io import BytesIO

import streamlit as st

try:
    from google import genai
    from google.genai import types
except ImportError:
    genai = None
    types = None

from docx import Document
from pypdf import PdfReader

APP_TITLE = "Resume ATS Analyzer"
DEFAULT_MODEL = "gemini-2.5-flash"

st.set_page_config(page_title=APP_TITLE, page_icon="📄", layout="wide")


def extract_text(uploaded_file) -> str:
    """Extract text from PDF, DOCX, or TXT uploads."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages).strip()

    if name.endswith(".docx"):
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="replace").strip()

    raise ValueError("Unsupported file type. Please upload PDF, DOCX, or TXT.")


def local_resume_checks(text: str) -> dict:
    """Simple non-AI checks used as a sanity signal and fallback."""
    lower = text.lower()
    section_terms = {
        "contact": ["email", "phone", "linkedin"],
        "summary": ["summary", "profile", "objective"],
        "experience": ["experience", "employment", "work history"],
        "education": ["education", "degree", "university", "college"],
        "skills": ["skills", "technical skills", "competencies"],
    }
    sections = {k: any(term in lower for term in v) for k, v in section_terms.items()}
    bullets = len(re.findall(r"(^|\n)\s*[•●▪◦*-]\s+", text))
    word_count = len(re.findall(r"\b\w+\b", text))
    action_words = [
        "developed", "built", "created", "implemented", "managed", "led",
        "improved", "designed", "analyzed", "automated", "optimized", "delivered",
    ]
    action_hits = sum(lower.count(word) for word in action_words)
    has_numbers = bool(re.search(r"\b\d+(?:%|\+)?\b", text))
    present = sum(sections.values())
    score = 25 + present * 8 + min(bullets, 8) * 2 + min(action_hits, 8) * 2 + (5 if has_numbers else 0)
    return {
        "score": min(score, 100),
        "sections": sections,
        "word_count": word_count,
        "bullet_count": bullets,
        "has_metrics": has_numbers,
    }


def parse_ai_json(raw: str) -> dict:
    raw = raw.strip()
    if raw.startswith("```"):
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```$", "", raw)
    start = raw.find("{")
    end = raw.rfind("}")
    if start == -1 or end == -1:
        raise ValueError("Gemini did not return valid JSON.")
    return json.loads(raw[start:end + 1])


def analyze_with_gemini(resume_text: str, job_description: str, api_key: str, model: str) -> dict:
    if genai is None:
        raise RuntimeError("google-genai is not installed. Install requirements.txt first.")

    client = genai.Client(api_key=api_key)
    jd_context = job_description.strip() or "No job description was supplied. Score ATS readiness using general recruiter/ATS best practices."
    prompt = f"""
You are an ATS resume reviewer. Analyze the resume below.

Return ONLY one JSON object with these keys:
score: integer 0-100
summary: short string
strengths: array of 3-5 strings
improvements: array of 5-8 strings
missing_keywords: array of strings
formatting_issues: array of strings
section_feedback: object with keys contact, summary, experience, education, skills
rewritten_summary: a concise improved professional summary; do not invent facts

Rules:
- If a job description is provided, prioritize its skills and terminology.
- Do not invent experience, education, certifications, metrics, employers, or technologies.
- Reward clear sections, measurable achievements, relevant keywords, action verbs, and ATS-readable formatting.
- Penalize vague wording, missing contact details, missing measurable impact, overly decorative formatting, and keyword stuffing.
- The score is a practical ATS-readiness estimate, not a guarantee of passing a real ATS.

JOB DESCRIPTION:
{jd_context}

RESUME:
{resume_text[:50000]}
"""

    response = client.models.generate_content(
        model=model,
        contents=prompt,
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return parse_ai_json(response.text)


def main():
    st.title("📄 Resume ATS Analyzer")
    st.caption("Upload a resume to get an ATS-readiness score and practical improvement suggestions.")

    with st.sidebar:
        st.header("Settings")
        api_key = st.text_input(
            "Gemini API key",
            value=st.secrets.get("GEMINI_API_KEY", os.getenv("GEMINI_API_KEY", "")),
            type="password",
            help="For deployment, store this in Streamlit Secrets instead of hard-coding it.",
        )
        model = st.text_input("Gemini model", value=DEFAULT_MODEL)
        st.info("Tip: adding a target job description makes the ATS score more job-specific.")

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Optional target job description",
        height=180,
        placeholder="Paste the job description here for keyword and relevance matching.",
    )

    if not uploaded:
        st.markdown("### How it works")
        st.markdown("1. Upload PDF, DOCX, or TXT.\n2. Optionally paste a job description.\n3. Click **Analyze Resume**.\n4. Review the score, strengths, and improvements.")
        return

    try:
        resume_text = extract_text(uploaded)
    except Exception as exc:
        st.error(f"Could not read the resume: {exc}")
        return

    if not resume_text:
        st.error("No readable text was found in the uploaded file. A scanned PDF may need OCR before analysis.")
        return

    with st.expander("Preview extracted resume text"):
        st.text(resume_text[:10000])

    if st.button("Analyze Resume", type="primary", use_container_width=True):
        local = local_resume_checks(resume_text)
        if not api_key:
            st.warning("No Gemini API key was supplied. Showing the local structural check only.")
            result = {
                "score": local["score"],
                "summary": "This is a local structural estimate. Add a Gemini API key for AI-based keyword and content review.",
                "strengths": ["Readable text was extracted successfully.", "The resume has detectable standard sections."],
                "improvements": ["Add a target job description for job-specific ATS matching.", "Use measurable achievements where possible.", "Use clear standard section headings."],
                "missing_keywords": [],
                "formatting_issues": [],
                "section_feedback": {k: ("Detected" if v else "Not clearly detected") for k, v in local["sections"].items()},
                "rewritten_summary": "",
            }
        else:
            with st.spinner("Analyzing your resume with Gemini..."):
                try:
                    result = analyze_with_gemini(resume_text, job_description, api_key, model)
                except Exception as exc:
                    st.error(f"Gemini analysis failed: {exc}")
                    st.stop()

        score = max(0, min(100, int(result.get("score", local["score"]))))
        st.divider()
        col1, col2 = st.columns([1, 2])
        with col1:
            st.metric("ATS Score", f"{score}/100")
            st.progress(score / 100)
        with col2:
            st.subheader("Summary")
            st.write(result.get("summary", "No summary returned."))

        st.subheader("Strengths")
        for item in result.get("strengths", []):
            st.success(item)

        st.subheader("Improvements")
        for item in result.get("improvements", []):
            st.warning(item)

        left, right = st.columns(2)
        with left:
            st.subheader("Missing / useful keywords")
            keywords = result.get("missing_keywords", [])
            st.write(", ".join(keywords) if keywords else "No major keyword gaps reported.")
        with right:
            st.subheader("Formatting issues")
            issues = result.get("formatting_issues", [])
            st.write("\n".join(f"- {x}" for x in issues) if issues else "No major formatting issues reported.")

        st.subheader("Section feedback")
        feedback = result.get("section_feedback", {})
        for section, note in feedback.items():
            st.markdown(f"**{section.title()}:** {note}")

        rewritten = result.get("rewritten_summary", "")
        if rewritten:
            st.subheader("Suggested summary")
            st.write(rewritten)

        st.caption("ATS scores are estimates. Different employers and ATS products use different rules.")


if __name__ == "__main__":
    main()
